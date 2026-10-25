"""Read XTB exports without relying on worksheet dimensions or row positions."""

import hashlib
import json
import re
import warnings
from collections import defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from app.models.enums import Currency

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 50000
TRADE_TYPES = {"Stock purchase": "buy", "Stock sell": "sell"}
UNIT_ADJUSTMENT = "Stock unit adjustment"
POSITION_TYPES = {**TRADE_TYPES, UNIT_ADJUSTMENT: "buy"}
CASH_TYPES = {
    "Deposit": ("Wpłata XTB", True),
    "Withdrawal": ("Wypłata XTB", True),
    "IKE deposit": ("Transfer IKE", True),
    "Subaccount transfer": ("Transfer wewnętrzny XTB", False),
    "Free funds interest": ("Odsetki XTB", False),
    "Free funds interest tax": ("Podatek od odsetek XTB", False),
    "Dividend": ("Dywidenda XTB", False),
    "Dividend tax": ("Podatek od dywidendy XTB", False),
}
EXECUTION = re.compile(
    r"^(OPEN|CLOSE) BUY (?P<volume>\d+(?:\.\d+)?)(?:/\d+(?:\.\d+)?)?"
    r" @ (?P<price>\d+(?:\.\d+)?)$"
)


def number(value, places: int = 8, positive: bool = False) -> Decimal:
    result = Decimal(str(value))
    if not result.is_finite() or abs(result) >= Decimal("10000000000"):
        raise ValueError("Liczba poza obsługiwanym zakresem")
    rounded = result.quantize(Decimal(1).scaleb(-places))
    if result != rounded or (positive and result <= 0):
        raise ValueError("Nieprawidłowa precyzja lub znak liczby")
    return rounded


def identifier(value) -> str:
    result = Decimal(str(value))
    if not result.is_finite() or result != result.to_integral_value() or result <= 0:
        raise ValueError("Nieprawidłowy identyfikator XTB")
    return str(int(result))


def timestamp(value, epoch) -> datetime:
    if isinstance(value, (float, int)):
        value = from_excel(value, epoch)
    if not isinstance(value, datetime) or value.date() > date.today():
        raise ValueError("Nieprawidłowa data operacji XTB")
    return value.replace(tzinfo=timezone.utc)


def fingerprint(event: dict) -> str:
    # File names, worksheet positions and report ranges never identify operations.
    source = {k: v for k, v in event.items() if k != "fingerprint"}
    return hashlib.sha256(
        json.dumps(
            source, sort_keys=True, ensure_ascii=False, separators=(",", ":")
        ).encode()
    ).hexdigest()


def sheet_rows(sheet):
    # XTB declares A1:A1 even when the worksheet contains hundreds of rows.
    sheet.reset_dimensions()
    for index, row in enumerate(sheet.iter_rows(values_only=True)):
        if index >= MAX_ROWS or len(row) > 100:
            raise ValueError("Raport przekracza limit 50000 wierszy lub 100 kolumn")
        yield row


def records(rows, required: set[str]):
    header = None
    for row in rows:
        if required.issubset({str(v) for v in row if v is not None}):
            header = [str(v) for v in row]
        elif header is not None:
            yield dict(zip(header, row))
    if header is None:
        raise ValueError("Brak wymaganych kolumn: " + ", ".join(sorted(required)))


def parse_file(filename: str, content: bytes, known_events: dict | None = None) -> dict:
    if not filename.lower().endswith(".xlsx") or len(content) > MAX_FILE_BYTES:
        raise HTTPException(422, "Dodaj plik XLSX o rozmiarze do 10 MB.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 80 * 1024 * 1024:
                raise ValueError("Zbyt duży rozpakowany raport")
        with warnings.catch_warnings():
            warnings.filterwarnings(
                "ignore", message="Workbook contains no default style"
            )
            workbook = load_workbook(
                BytesIO(content), read_only=True, data_only=True, keep_links=False
            )
        try:
            return _parse_workbook(filename, workbook, known_events or {})
        finally:
            workbook.close()
    except (
        ValueError,
        InvalidOperation,
        BadZipFile,
        KeyError,
        TypeError,
        OSError,
        StopIteration,
        ParseError,
    ) as exc:
        raise HTTPException(
            422, f"{filename}: nieprawidłowy raport XTB. {exc}"
        ) from exc


def _parse_workbook(filename, workbook, known_events) -> dict:
    required = {"Cash Operations", "Closed Positions", "Open Positions"}
    if not required.issubset(workbook.sheetnames):
        raise ValueError("Potrzebny pełny raport z trzema arkuszami XTB")
    cash_rows = list(sheet_rows(workbook["Cash Operations"]))
    account = identifier(
        next(
            row[1] for row in cash_rows if len(row) >= 2 and row[0] == "Account number"
        )
    )
    notes = []
    # Currency is present in position summaries, except in the supplied IKE export.
    open_rows = list(sheet_rows(workbook["Open Positions"]))
    currency_values = {
        row[3]
        for row in open_rows
        if len(row) >= 4
        and row[1] in ("Open position value", "Open position profit")
        and row[3] in Currency.__members__
    }
    prefix = filename.split("_", 1)[0].upper()
    is_ike = prefix == "IKE" or any(row and row[0] == "IKE" for row in open_rows)
    currency = next(iter(currency_values)) if len(currency_values) == 1 else None
    if is_ike:
        currency = "PLN"
    elif currency is None and prefix in Currency.__members__:
        currency = prefix
        notes.append(
            f"{filename}: waluta {currency} pochodzi z nazwy pliku. Sprawdź ją."
        )
    if not is_ike and (
        len(currency_values) > 1
        or (currency_values and prefix in Currency.__members__ and prefix != currency)
    ):
        currency = None
        notes.append(f"{filename}: sprzeczne waluty. Wybierz walutę rachunku.")
    for row in open_rows:
        if (
            not is_ike
            and len(row) >= 2
            and row[0] == "Account number"
            and str(row[1]) != account
        ):
            notes.append(
                f"{filename}: numer rachunku w Open Positions różni się od historii. "
                f"Import używa rachunku {account} z Cash Operations."
            )
    instruments: dict[str, dict] = {}
    expected = defaultdict(lambda: Decimal(0))
    for row in records(
        open_rows, {"Ticker", "Category", "Volume", "Instrument/Position"}
    ):
        if row.get("Ticker") and row.get("Category") in ("ETF", "STOCK"):
            ticker = str(row["Ticker"]).upper()
            expected[ticker] += number(row["Volume"], positive=True)
            instruments[ticker] = {
                "name": str(row["Instrument/Position"]),
                "category": row["Category"],
                "issues": [],
            }
    as_of = timestamp(
        next(
            row[1]
            for row in open_rows
            if len(row) >= 2 and row[0] == "Data as of report generated"
        ),
        workbook.epoch,
    ).isoformat()
    events = []
    for row in records(
        cash_rows, {"Type", "Time", "Amount", "ID", "Comment", "Ticker"}
    ):
        if row.get("Type") in (None, "", "Total"):
            continue
        kind = row["Type"]
        if kind not in TRADE_TYPES and kind not in CASH_TYPES:
            raise ValueError(f"Nieobsługiwany typ operacji: {kind}")
        amount = number(row["Amount"], 2)
        ticker = str(row.get("Ticker") or "").upper()
        event = {
            "external_id": f"xtb:{account}:cash:{identifier(row['ID'])}",
            "account": account,
            "kind": kind,
            "amount": str(amount),
            "at": timestamp(row["Time"], workbook.epoch).isoformat(),
            "ticker": ticker,
            "comment": str(row.get("Comment") or ""),
            "position_id": identifier(row["Position ID"])
            if row.get("Position ID")
            else None,
            "product": str(row.get("Product") or ""),
        }
        if kind in TRADE_TYPES:
            if not ticker or len(ticker) > 64:
                raise ValueError("Brak prawidłowego tickera transakcji")
            instrument = instruments.setdefault(
                ticker,
                {
                    "name": str(row.get("Instrument") or ticker),
                    "category": str(row.get("Category") or ""),
                    "issues": [],
                },
            )
            match = EXECUTION.fullmatch(event["comment"].strip())
            if match is None:
                instrument["issues"].append("Nierozpoznany opis wykonania transakcji.")
            else:
                event["volume"] = str(number(match["volume"], positive=True))
                if (kind == "Stock purchase") != (match[1] == "OPEN"):
                    raise ValueError("Typ transakcji nie zgadza się z opisem wykonania")
            if amount == 0 or (kind == "Stock purchase") != (amount < 0):
                raise ValueError("Nieprawidłowy znak kwoty transakcji")
            if instrument["category"] not in ("ETF", "STOCK"):
                instrument["issues"].append(
                    "Obsługiwane są akcje i ETF; bez CFD i krótkiej sprzedaży."
                )
        elif ticker:
            # Dividends may refer to a stock without creating a position.
            event["ticker"] = ""
        event["fingerprint"] = fingerprint(event)
        events.append(event)

    closed = list(
        records(
            sheet_rows(workbook["Closed Positions"]),
            {"Position ID", "Ticker", "Close Origin"},
        )
    )
    lots = replacement_lots(open_rows, closed, workbook.epoch)
    source_events = {**known_events, **{e["external_id"]: e for e in events}}
    sales = {
        e["position_id"] for e in source_events.values() if e["kind"] == "Stock sell"
    }
    for row in closed:
        if not row.get("Position ID"):
            continue
        ticker = str(row.get("Ticker") or "").upper()
        position_id = identifier(row["Position ID"])
        instrument = instruments.get(ticker)
        if instrument is not None and (
            row.get("Close Origin") == "Correction"
            or "STC" in str(row.get("Comment", ""))
        ):
            adjustment = unit_adjustment(
                account, row, lots, source_events, workbook.epoch
            )
            if adjustment is None:
                instrument["issues"].append(
                    "Nie można jednoznacznie powiązać korekty STC z nową pozycją. "
                    "Dodaj raport obejmujący obie pozycje lub pomiń instrument."
                )
            else:
                events.append(adjustment)
                notes.append(
                    f"{ticker}: przeliczenie jednostek {adjustment['ratio']}:1 "
                    "z zachowaniem kosztu zakupu i bez przepływu gotówki."
                )
        elif position_id not in sales:
            raise ValueError(
                f"Brak operacji gotówkowej dla zamkniętej pozycji {position_id}"
            )

    for instrument in instruments.values():
        instrument["issues"] = sorted(set(instrument["issues"]))
    return {
        "account": account,
        "filename": filename,
        "currency": currency,
        "is_ike": is_ike,
        "events": events,
        "instruments": instruments,
        "warnings": notes,
        "snapshot": {
            "at": as_of,
            "positions": {k: str(v) for k, v in expected.items()},
        },
    }


def replacement_lots(open_rows, closed, epoch) -> list[dict]:
    """Reconstruct original replacement lots, including any subsequently sold part."""
    lots: dict[str, dict] = {}
    for row in records(open_rows, {"Instrument/Position", "Type", "Open price"}):
        if row.get("Type") == "BUY" and row.get("Ticker"):
            key = identifier(row["Instrument/Position"])
            lots[key] = {
                "id": key,
                "ticker": str(row["Ticker"]).upper(),
                "volume": number(row["Volume"], positive=True),
                "price": number(row["Open price"], positive=True),
                "at": timestamp(row["Open time (UTC)"], epoch),
            }
    for row in closed:
        if not row.get("Position ID") or row.get("Type") != "BUY":
            continue
        key = identifier(row["Position ID"])
        lot = lots.setdefault(
            key,
            {
                "id": key,
                "ticker": str(row["Ticker"]).upper(),
                "volume": Decimal(0),
                "price": number(row["Open Price"], positive=True),
                "at": timestamp(row["Open Time (UTC)"], epoch),
            },
        )
        lot["volume"] += number(row["Volume"], positive=True)
    return list(lots.values())


def unit_adjustment(account, row, lots, source_events, epoch) -> dict | None:
    """Recognize only a proven forward unit conversion, never guess a transfer."""
    if row.get("Comment") != "STC Transfer Out" or row.get("Type") != "BUY":
        return None
    old_id = identifier(row["Position ID"])
    closed_at = timestamp(row["Close Time (UTC)"], epoch)
    old_volume = number(row["Volume"], positive=True)
    old_price = number(row["Open Price"], positive=True)
    opened_at = timestamp(row["Open Time (UTC)"], epoch)
    ticker = str(row["Ticker"]).upper()
    external_id = f"xtb:{account}:stc:{old_id}:{closed_at.isoformat()}"
    correction = {
        "position_id": old_id,
        "volume": str(old_volume),
        "price": str(old_price),
        "opened_at": opened_at.isoformat(),
        "closed_at": closed_at.isoformat(),
    }
    known = source_events.get(external_id)
    if known is not None:
        # Later exports can have a smaller remaining lot, or none after a sale.
        # Reuse confirmed evidence, while rejecting changed correction details.
        return known if known.get("correction") == correction else None
    original = sum(
        (
            Decimal(e["volume"]) * (1 if e["kind"] == "Stock purchase" else -1)
            for e in source_events.values()
            if e["account"] == account
            and e.get("position_id") == old_id
            and e["kind"] in TRADE_TYPES
            and "volume" in e
            and e["at"] <= closed_at.isoformat()
        ),
        Decimal(0),
    )
    if original != old_volume:
        return None
    matches = []
    for lot in lots:
        ratio = lot["volume"] / old_volume
        if (
            lot["id"] != old_id
            and lot["ticker"] == ticker
            and abs((lot["at"] - opened_at).total_seconds()) <= 36 * 3600
            and 1 < ratio <= 10000
            and ratio == ratio.to_integral_value()
            and old_price == lot["price"] * ratio
        ):
            matches.append((lot, ratio))
    if len(matches) != 1:
        return None
    lot, ratio = matches[0]
    event = {
        "external_id": external_id,
        "account": account,
        "kind": UNIT_ADJUSTMENT,
        "amount": "0.00",
        "at": closed_at.isoformat(),
        "ticker": ticker,
        "volume": str(number(lot["volume"] - old_volume, positive=True)),
        "position_id": old_id,
        "replacement_position_id": lot["id"],
        "ratio": str(int(ratio)),
        "correction": correction,
        "comment": "STC Transfer Out",
        "product": str(row.get("Product") or ""),
    }
    event["fingerprint"] = fingerprint(event)
    return event


def parse_files(
    files: list[tuple[str, bytes]], known_events: dict | None = None
) -> dict:
    if len(files) != 1:
        raise HTTPException(422, "Dodaj jeden plik XLSX naraz.")
    accounts, instruments, events, notes, snapshots = [], {}, {}, [], {}
    currencies = defaultdict(set)
    for filename, content in files:
        report = parse_file(filename, content, {**(known_events or {}), **events})
        accounts.append(
            {k: report[k] for k in ("account", "filename", "currency", "is_ike")}
        )
        accounts[-1]["operation_count"] = len(report["events"])
        if report["currency"]:
            currencies[report["account"]].add(report["currency"])
        notes.extend(report["warnings"])
        previous_snapshot = snapshots.get(report["account"])
        if (
            previous_snapshot is None
            or previous_snapshot["at"] < report["snapshot"]["at"]
        ):
            snapshots[report["account"]] = report["snapshot"]
        for ticker, info in report["instruments"].items():
            if ticker in instruments:
                info["issues"] = sorted(
                    set(info["issues"] + instruments[ticker]["issues"])
                )
            instruments[ticker] = info
        for event in report["events"]:
            previous = events.get(event["external_id"])
            if previous and previous["fingerprint"] != event["fingerprint"]:
                raise HTTPException(
                    409, f"Sprzeczna treść operacji {event['external_id']} w plikach."
                )
            events[event["external_id"]] = event
    ike_accounts = {a["account"] for a in accounts if a["is_ike"]}
    for account in accounts:
        if account["account"] in ike_accounts:
            account["currency"], account["is_ike"] = "PLN", True
            continue
        if len(currencies[account["account"]]) > 1:
            account["currency"] = None
            notes.append(
                f"Sprzeczne waluty rachunku {account['account']}. Wybierz właściwą."
            )
    return {
        "accounts": accounts,
        "instruments": instruments,
        "events": sorted(events.values(), key=lambda e: (e["at"], e["external_id"])),
        "warnings": sorted(set(notes)),
        "snapshots": snapshots,
    }
